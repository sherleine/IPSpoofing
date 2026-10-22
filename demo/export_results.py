"""
Exports clutch_results.json to Excel and CSV.
Opens the Excel file automatically when done.
"""
import json
import csv
import os
import subprocess
from datetime import datetime

INPUT  = "output/clutch_results.json"
STAMP  = datetime.now().strftime("%Y-%m-%d_%H-%M")
XLSX   = f"output/clutch_results_{STAMP}.xlsx"
CSV    = f"clutch_results_{STAMP}.csv"

# Load data
with open(INPUT, encoding="utf-8") as f:
    companies = json.load(f)

# Add rank column
for i, c in enumerate(companies, 1):
    c["rank"] = i

FIELDS = ["rank", "name", "rating", "reviews", "location"]

# ── CSV ──────────────────────────────────────────────────
with open(CSV, "w", newline="", encoding="utf-8-sig") as f:
    writer = csv.DictWriter(f, fieldnames=FIELDS, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(companies)
print(f"CSV saved : {CSV}")

# ── Excel ─────────────────────────────────────────────────
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clutch Top Companies"

    # Header row
    headers = ["#", "Company", "Rating", "Reviews", "Location"]
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", size=11)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Data rows
    for row_i, c in enumerate(companies, 2):
        ws.cell(row=row_i, column=1, value=c["rank"])
        ws.cell(row=row_i, column=2, value=c["name"])
        ws.cell(row=row_i, column=3, value=c["rating"])
        ws.cell(row=row_i, column=4, value=c["reviews"])
        ws.cell(row=row_i, column=5, value=c["location"])

        # Alternate row colour
        if row_i % 2 == 0:
            fill = PatternFill("solid", fgColor="D6E4F0")
            for col in range(1, 6):
                ws.cell(row=row_i, column=col).fill = fill

    # Column widths
    widths = [5, 42, 8, 10, 28]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Freeze header row
    ws.freeze_panes = "A2"

    wb.save(XLSX)
    print(f"Excel saved: {XLSX}")

    # Open Excel file automatically
    os.startfile(os.path.abspath(XLSX))

except ImportError:
    print("openpyxl not installed — install it with: pip install openpyxl")
    print(f"CSV is ready at: {CSV}")
